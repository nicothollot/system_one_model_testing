"""Input preservation, PDF preprocessing, simulation, and lossless exports."""
import csv
import hashlib
import io
import json
import math
from pathlib import Path

from app.criteria import QUESTIONS, OPTIONS
TEMPLATE = "EXTRACTION OBJECTIVE\n{objective}\n\nPAGE {page}\n{text}"


def parse_json(raw):
    if len(raw) > 2 * 1024 * 1024:
        raise ValueError("Instructions JSON exceeds the prototype's 2 MiB input limit; no content was truncated")
    def invalid(value):
        raise ValueError(f"Nonfinite JSON number: {value}")
    try:
        value = json.loads(raw, parse_constant=invalid)
    except (ValueError, UnicodeError) as exc:
        raise ValueError(f"Invalid instructions JSON: {exc}") from exc
    if not isinstance(value, (dict, list)) or not value:
        raise ValueError("Instructions JSON must be a nonempty object or array")
    return value


def parse_xlsx(raw):
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        cached = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheets = []
        cell_count = 0
        character_count = 0
        try:
            for sheet in workbook:
                rows = []
                for cells, values in zip(sheet.iter_rows(), cached[sheet.title].iter_rows()):
                    populated = []
                    for cell, value in zip(cells, values):
                        if cell.value is not None:
                            cell_count += 1
                            character_count += len(str(cell.value))
                            if cell_count > 100000 or character_count > 2 * 1024 * 1024:
                                raise ValueError("Workbook exceeds 100,000 populated cells or 2 MiB of reference text; no content was truncated")
                            item = {"cell": cell.coordinate, "value": str(cell.value)}
                            if cell.data_type == "f":
                                item["cached_value"] = None if value.value is None else str(value.value)
                            populated.append(item)
                    if populated:
                        rows.append(populated)
                sheets.append({"sheet": sheet.title, "rows": rows})
        finally:
            workbook.close()
            cached.close()
    except Exception as exc:
        raise ValueError(f"Invalid reference XLSX: {exc}") from exc
    if not any(s["rows"] for s in sheets):
        raise ValueError("Reference XLSX contains no populated cells")
    return sheets


def objective(instructions, reference):
    from app.objective import compile_objective
    return compile_objective(instructions, reference)["compiled_routing_objective"]


def preprocess(raw):
    import pymupdf
    pages = []
    try:
        with pymupdf.open(stream=raw, filetype="pdf") as pdf:
            if pdf.needs_pass:
                raise ValueError("Encrypted PDF requires a decrypted copy")
            for number in range(len(pdf)):
                error = None
                try:
                    text = pdf[number].get_text("text", sort=True)
                except Exception as exc:
                    text, error = "", f"PyMuPDF extraction failed: {exc}"
                sparse = len(text.strip()) < 80
                warning = error or ("Little/no extractable text; visual or OCR review required" if sparse else None)
                pages.append({"page": number + 1, "extracted_text": text, "characters": len(text),
                              "approx_tokens": math.ceil(len(text) / 4), "text_extraction_status": "error" if error else "sparse" if sparse else "ok",
                              "text_extraction_warning": warning, "needs_visual_or_ocr_review": sparse or bool(error)})
    except Exception as exc:
        raise ValueError(f"PDF preprocessing failed: {exc}") from exc
    if not pages:
        raise ValueError("PDF has no pages")
    return pages


def validate_scores(scores):
    if set(scores) != set(QUESTIONS):
        raise ValueError("Missing classifier distributions")
    for distribution in scores.values():
        if set(distribution) != {"yes", "no"}:
            raise ValueError("Expected yes/no distribution")
        if any(not math.isfinite(p) or not 0 <= p <= 1 for p in distribution.values()):
            raise ValueError("Invalid probability")
        if not math.isclose(sum(distribution.values()), 1.0, abs_tol=1e-9):
            raise ValueError("Probabilities do not sum to one")


def select(pages, threshold):
    if not 0 <= threshold <= 1:
        raise ValueError("Threshold must be in [0, 1]")
    return [p["page"] for p in pages if p.get("scores") and p["scores"]["overall_relevance"]["yes"] > threshold]


def expand(selected, count):
    return sorted({n + delta for n in selected for delta in (-1, 0, 1) if 1 <= n + delta <= count})


def thresholds(pages):
    if any(not p.get("scores") for p in pages):
        return {}  # Unknown pages cannot support document-wide reduction statistics.
    result = {}
    for threshold in (.10, .25, .50, .75, .90, .95):
        chosen = select(pages, threshold)
        result[f"{threshold:.2f}"] = {"pages": chosen, "count": len(chosen), "retained_percent": 100 * len(chosen) / len(pages)}
    return result


def csv_rows(pages):
    result = []
    for page in pages:
        row = {"page": page["page"]}
        for key in QUESTIONS:
            for option in ("yes", "no"):
                row[f"{key}_{option}"] = (page.get("scores") or {}).get(key, {}).get(option)
        row.update(characters=page["characters"], tokens=page.get("tokens", page["approx_tokens"]),
                   routing_objective_tokens=page.get("routing_objective_tokens"), page_text_tokens=page.get("page_text_tokens"),
                   exact_prompt_tokens_min=min(page.get("exact_prompt_tokens", {}).values(), default=None),
                   exact_prompt_tokens_max=max(page.get("exact_prompt_tokens", {}).values(), default=None),
                   available_token_headroom=min(page.get("available_token_headroom", {}).values(), default=None),
                   classification_ms=1000 * page.get("classification_seconds", 0),
                   needs_visual_or_ocr_review=page["needs_visual_or_ocr_review"],
                   needs_router_review=page.get("needs_router_review", not bool(page.get("scores"))),
                   classification_error=page.get("classification_error", page.get("exception")),
                   status=page.get("classification_status"), exception=page.get("exception"),
                   text_warning=page["text_extraction_warning"])
        result.append(row)
    return result


def export(result, folder, stem):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    paths = {suffix: str(folder / f"{stem}.{suffix}") for suffix in ("json", "csv")}
    result["run"]["exports"] = paths
    rows = csv_rows(result["pages"])
    temporary = Path(paths["csv"] + ".tmp")
    with temporary.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(paths["csv"])
    temporary = Path(paths["json"] + ".tmp")
    temporary.write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    temporary.replace(paths["json"])
    return paths


def fingerprint(raw):
    return hashlib.sha256(raw).hexdigest()
